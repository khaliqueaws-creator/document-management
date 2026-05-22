from pathlib import Path
import math
import textwrap

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageDraw, ImageFont, ImageOps


BASE_DIR = Path(__file__).resolve().parent
WORD_DIR = BASE_DIR / "word"
EXCEL_DIR = BASE_DIR / "excel"
PDF_DIR = BASE_DIR / "pdf"
OCR_DIR = BASE_DIR / "ocr_images"


DOCUMENTS = [
    ("employee_onboarding_policy", "HR", "Policy", "Employee Onboarding Policy", "New hire orientation, identity verification, equipment pickup, security badge issuance, payroll setup, benefits enrollment, and first-week training checklist."),
    ("benefits_enrollment_summary", "HR", "Summary", "Benefits Enrollment Summary", "Medical insurance, dental coverage, retirement contributions, dependent verification, open enrollment deadlines, and employee support contacts."),
    ("remote_work_guidelines", "HR", "Policy", "Remote Work Guidelines", "Hybrid schedule eligibility, home office expectations, secure network access, manager approval, and productivity reporting requirements."),
    ("vendor_invoice_cloud_services", "Finance", "Invoice", "Vendor Invoice - Cloud Services", "Monthly invoice for cloud hosting, Kubernetes platform support, storage usage, managed database fees, and payment terms."),
    ("purchase_order_laptops", "Finance", "Purchase Order", "Laptop Purchase Order", "Approved procurement request for developer laptops, docking stations, warranty coverage, cost center allocation, and delivery schedule."),
    ("expense_reimbursement_policy", "Finance", "Policy", "Expense Reimbursement Policy", "Travel receipts, meal limits, mileage reimbursement, approval workflow, submission deadlines, and audit documentation requirements."),
    ("quarterly_budget_review", "Finance", "Report", "Quarterly Budget Review", "Department spending, forecast variance, capital expense tracking, operating budget risks, and corrective action recommendations."),
    ("password_reset_procedure", "IT", "Procedure", "Password Reset Procedure", "Identity verification, temporary password issuance, multifactor authentication reset, service desk ticket notes, and account recovery steps."),
    ("security_access_request", "IT", "Request", "Security Access Request", "Role-based access approval, privileged account review, application entitlements, manager signoff, and access expiration date."),
    ("data_retention_policy", "IT", "Policy", "Data Retention Policy", "Document retention schedules, encrypted archive storage, legal hold requirements, deletion approvals, and compliance reporting."),
    ("incident_response_report", "IT", "Report", "Incident Response Report", "Security event timeline, affected systems, containment steps, root cause, remediation plan, and lessons learned."),
    ("software_release_notes", "IT", "Release Notes", "Software Release Notes", "Application version update, bug fixes, deployment window, rollback procedure, known issues, and support contacts."),
    ("service_level_agreement", "Legal", "Contract", "Service Level Agreement", "Availability targets, support response times, service credits, maintenance windows, escalation contacts, and reporting obligations."),
    ("contract_termination_notice", "Legal", "Notice", "Contract Termination Notice", "Termination effective date, notice period, outstanding obligations, final invoice handling, and return of confidential materials."),
    ("mutual_nda_template", "Legal", "Agreement", "Mutual NDA Template", "Confidential information, permitted disclosures, exclusions, term duration, return of materials, and governing law."),
    ("facilities_maintenance_request", "Operations", "Request", "Facilities Maintenance Request", "Office repair ticket, priority level, room location, vendor dispatch, safety concern, and completion notes."),
    ("warehouse_inventory_audit", "Operations", "Audit", "Warehouse Inventory Audit", "Cycle count results, inventory variance, damaged items, reconciliation actions, and supervisor approval."),
    ("project_status_modernization", "Operations", "Report", "Project Status - Modernization", "Milestone progress, schedule risks, dependency tracking, budget status, stakeholder decisions, and next steps."),
    ("training_attendance_roster", "Operations", "Roster", "Training Attendance Roster", "Employee training session, attendance confirmation, course completion, instructor notes, and follow-up assignments."),
    ("customer_complaint_summary", "Operations", "Summary", "Customer Complaint Summary", "Complaint category, customer impact, investigation findings, resolution owner, and preventive action."),
    ("quality_control_checklist", "Operations", "Checklist", "Quality Control Checklist", "Inspection criteria, acceptance thresholds, defect tracking, reviewer initials, and release approval."),
    ("travel_authorization_form", "Finance", "Form", "Travel Authorization Form", "Business travel purpose, destination, estimated airfare, hotel budget, manager approval, and policy acknowledgement."),
    ("equipment_return_notice", "HR", "Notice", "Equipment Return Notice", "Employee offboarding, laptop return, badge deactivation, mobile device collection, and final inventory confirmation."),
    ("compliance_training_plan", "HR", "Plan", "Compliance Training Plan", "Mandatory training calendar, privacy awareness, workplace conduct, completion tracking, and escalation for overdue staff."),
    ("database_backup_procedure", "IT", "Procedure", "Database Backup Procedure", "Backup schedule, retention policy, restore testing, encrypted storage, monitoring alerts, and disaster recovery owner."),
    ("api_access_review", "IT", "Review", "API Access Review", "Service account inventory, token expiration, unused permissions, integration owner, and remediation tracking."),
    ("supplier_risk_assessment", "Legal", "Assessment", "Supplier Risk Assessment", "Vendor risk score, data processing review, contract safeguards, insurance coverage, and approval conditions."),
    ("board_meeting_minutes", "Operations", "Minutes", "Board Meeting Minutes", "Agenda items, approvals, budget discussion, policy updates, assigned actions, and next meeting date."),
    ("marketing_campaign_brief", "Marketing", "Brief", "Marketing Campaign Brief", "Audience segment, campaign goals, launch timeline, creative assets, budget, and success metrics."),
    ("customer_case_study", "Marketing", "Case Study", "Customer Case Study", "Business challenge, implemented solution, measurable outcomes, customer quote, and publishing approval."),
    ("sales_pipeline_report", "Sales", "Report", "Sales Pipeline Report", "Opportunity stages, forecast amount, close probability, account owner, next action, and quarter target."),
    ("renewal_quote_summary", "Sales", "Quote", "Renewal Quote Summary", "Subscription renewal, license count, discount approval, payment terms, and customer acceptance deadline."),
    ("privacy_impact_assessment", "Legal", "Assessment", "Privacy Impact Assessment", "Personal data categories, processing purpose, retention period, access controls, and privacy risk mitigation."),
    ("employee_performance_review", "HR", "Review", "Employee Performance Review", "Performance goals, manager feedback, development plan, promotion readiness, and acknowledgement signature."),
    ("network_change_request", "IT", "Request", "Network Change Request", "Firewall rule update, affected subnet, change window, test plan, rollback steps, and approval history."),
    ("asset_inventory_register", "Operations", "Register", "Asset Inventory Register", "Asset tag, owner department, assigned employee, purchase date, device condition, and replacement cycle."),
    ("invoice_dispute_memo", "Finance", "Memo", "Invoice Dispute Memo", "Billing discrepancy, purchase order reference, vendor response, disputed amount, and resolution timeline."),
    ("policy_exception_request", "IT", "Request", "Policy Exception Request", "Exception justification, compensating controls, approval expiry, business owner, and security review notes."),
    ("office_relocation_plan", "Operations", "Plan", "Office Relocation Plan", "Move schedule, seating chart, network readiness, vendor coordination, employee communications, and risk checklist."),
    ("new_vendor_onboarding", "Finance", "Procedure", "New Vendor Onboarding", "Tax documentation, banking validation, supplier contact, purchase approval, compliance screening, and payment setup."),
]


OCR_DOCUMENTS = [
    ("ocr_employee_onboarding_notice.png", "Employee Onboarding Notice", "Bring identification, complete payroll forms, pick up laptop, and attend security training."),
    ("ocr_vendor_invoice_scan.jpg", "Vendor Invoice Scan", "Invoice 1048 for managed Kubernetes support, storage fees, and database monitoring services."),
    ("ocr_security_access_request.tiff", "Security Access Request", "Request privileged access for the reporting dashboard. Manager approval required before activation."),
    ("ocr_expense_receipt.jpeg", "Expense Receipt", "Travel meal reimbursement receipt for client meeting. Total amount 48.75 pending finance approval."),
    ("ocr_contract_notice.bmp", "Contract Notice", "Notice of contract renewal review. Legal team must confirm terms before the expiration date."),
    ("ocr_incident_report.png", "Incident Report", "User account lockout incident resolved by service desk after identity verification and MFA reset."),
    ("ocr_benefits_summary.jpg", "Benefits Summary", "Open enrollment includes health insurance, dental coverage, retirement plan, and dependent updates."),
    ("ocr_purchase_approval.tiff", "Purchase Approval", "Purchase request for ten laptops, docking stations, and warranty support for engineering team."),
    ("ocr_training_roster.jpeg", "Training Roster", "Compliance training attendance roster for privacy awareness and workplace conduct session."),
    ("ocr_inventory_count.bmp", "Inventory Count", "Warehouse cycle count found two damaged scanners and one missing barcode printer."),
]


def clean_name(name):
    return name.replace("_", " ").title()


def ensure_dirs():
    for directory in [WORD_DIR, EXCEL_DIR, PDF_DIR, OCR_DIR]:
        directory.mkdir(parents=True, exist_ok=True)


def make_paragraph(title, department, document_type, content):
    return (
        f"{title}\n\n"
        f"Department: {department}\n"
        f"Document Type: {document_type}\n\n"
        f"Purpose: This synthetic test document is designed for document "
        f"management, metadata extraction, OCR review, and semantic search.\n\n"
        f"Details: {content}\n\n"
        f"Review Notes: Use this file to test searches related to "
        f"{department.lower()}, {document_type.lower()}, approvals, policies, "
        f"reports, invoices, and operational workflows."
    )


def create_docx(slug, department, document_type, title, content):
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(f"Department: {department}")
    document.add_paragraph(f"Document Type: {document_type}")
    document.add_heading("Purpose", level=2)
    document.add_paragraph(
        "This synthetic test document supports upload, text extraction, "
        "AI metadata suggestions, embeddings, and semantic search testing."
    )
    document.add_heading("Details", level=2)
    for line in textwrap.wrap(content, width=86):
        document.add_paragraph(line)
    document.add_heading("Search Terms", level=2)
    document.add_paragraph(
        f"{department}, {document_type}, {clean_name(slug)}, approval, "
        "workflow, review, policy, report"
    )
    document.save(WORD_DIR / f"{slug}.docx")


def create_xlsx(slug, department, document_type, title, content):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Document Data"
    headers = ["Field", "Value"]
    sheet.append(headers)
    rows = [
        ("Title", title),
        ("Department", department),
        ("Document Type", document_type),
        ("Summary", content),
        ("Owner", f"{department} Team"),
        ("Status", "Ready for review"),
        ("Search Terms", f"{department}, {document_type}, {clean_name(slug)}"),
    ]
    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")

    sheet.column_dimensions["A"].width = 22
    sheet.column_dimensions["B"].width = 96
    workbook.save(EXCEL_DIR / f"{slug}.xlsx")


def pdf_escape(value):
    return value.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def create_pdf(slug, department, document_type, title, content):
    lines = []
    for block in make_paragraph(title, department, document_type, content).split("\n"):
        if not block:
            lines.append("")
            continue
        lines.extend(textwrap.wrap(block, width=86) or [""])

    text_commands = ["BT", "/F1 11 Tf", "72 740 Td", "14 TL"]
    for index, line in enumerate(lines[:45]):
        if index:
            text_commands.append("T*")
        text_commands.append(f"({pdf_escape(line)}) Tj")
    text_commands.append("ET")
    stream = "\n".join(text_commands).encode("latin-1", errors="replace")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream + b"\nendstream",
    ]

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode("ascii"))
        output.extend(body)
        output.extend(b"\nendobj\n")

    xref_offset = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )

    (PDF_DIR / f"{slug}.pdf").write_bytes(output)


def load_font(size=26):
    for font_name in ["arial.ttf", "calibri.ttf", "DejaVuSans.ttf"]:
        try:
            return ImageFont.truetype(font_name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def create_ocr_image(filename, title, body, rotation=0):
    width, height = 1250, 1650
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title_font = load_font(44)
    body_font = load_font(30)
    small_font = load_font(24)

    draw.rectangle((55, 55, width - 55, height - 55), outline=(80, 80, 80), width=3)
    draw.text((90, 95), title, fill=(20, 20, 20), font=title_font)
    draw.line((90, 160, width - 90, 160), fill=(80, 80, 80), width=2)
    draw.text((90, 205), "Synthetic OCR Test Document", fill=(40, 40, 40), font=small_font)
    draw.text((90, 255), "Department: Test Records", fill=(40, 40, 40), font=small_font)

    y = 340
    for line in textwrap.wrap(body, width=58):
        draw.text((90, y), line, fill=(10, 10, 10), font=body_font)
        y += 48

    y += 40
    table_rows = [
        ("Status", "Ready for Review"),
        ("Priority", "Normal"),
        ("Approval", "Required"),
        ("Search", title.lower()),
    ]
    for label, value in table_rows:
        draw.rectangle((90, y, width - 90, y + 58), outline=(120, 120, 120), width=2)
        draw.line((360, y, 360, y + 58), fill=(120, 120, 120), width=2)
        draw.text((110, y + 12), label, fill=(20, 20, 20), font=small_font)
        draw.text((385, y + 12), value, fill=(20, 20, 20), font=small_font)
        y += 58

    for stripe in range(0, height, 120):
        shade = 248 + int(4 * math.sin(stripe))
        draw.line((60, stripe, width - 60, stripe + 8), fill=(shade, shade, shade), width=1)

    image = ImageOps.autocontrast(image)
    if rotation:
        image = image.rotate(rotation, expand=True, fillcolor="white")

    suffix = Path(filename).suffix.lower()
    save_path = OCR_DIR / filename
    if suffix in [".jpg", ".jpeg"]:
        image.save(save_path, quality=92)
    else:
        image.save(save_path)


def main():
    ensure_dirs()

    for index, item in enumerate(DOCUMENTS):
        slug, department, document_type, title, content = item
        if index < 15:
            create_docx(slug, department, document_type, title, content)
        elif index < 30:
            create_xlsx(slug, department, document_type, title, content)
        else:
            create_pdf(slug, department, document_type, title, content)

    for index, item in enumerate(OCR_DOCUMENTS):
        filename, title, body = item
        rotation = [-1, 0, 1, -2, 2][index % 5]
        create_ocr_image(filename, title, body, rotation=rotation)

    total = len(DOCUMENTS) + len(OCR_DOCUMENTS)
    print(f"Generated {total} files in {BASE_DIR}")
    print(f"Word: {len(list(WORD_DIR.glob('*.docx')))}")
    print(f"Excel: {len(list(EXCEL_DIR.glob('*.xlsx')))}")
    print(f"PDF: {len(list(PDF_DIR.glob('*.pdf')))}")
    print(f"OCR images: {len(list(OCR_DIR.iterdir()))}")


if __name__ == "__main__":
    main()
