from django.core.files.storage import FileSystemStorage
from django.core.files import File
from django.http import FileResponse, Http404, HttpResponse
from django.shortcuts import render, redirect
from django.conf import settings
from django.contrib import messages
from django.contrib.auth import logout as django_logout
from django.core.exceptions import ValidationError
from django.core.paginator import Paginator
from django.db.models import Q
from django.urls import reverse
from django.utils import timezone
from uuid import uuid4

from .forms import (
    DocumentForm,
    DocumentMetadataForm,
    validate_uploaded_file,
)
from .ai_metadata import MetadataSuggestionError, suggest_metadata
from .embeddings import EmbeddingError, rebuild_document_embeddings
from .mcp_indexing import index_document as mcp_index_document
from .metadata_quality import MetadataQualityError, review_document_metadata
from .models import (
    AuditEvent,
    Document,
    DocumentChunk,
    MetadataQualityReview,
)
from .opensearch_indexing import (
    OpenSearchIndexingError,
    delete_document as delete_indexed_document,
    reindex_document,
)
from .rag import RAGError, answer_question
from .semantic_search import search_documents_by_meaning
from .auth import oauth
from .permissions import (
    okta_role_required,
    is_viewer,
    is_loader,
    is_admin,
)

import os
import json
import jwt
from pypdf import PdfReader
from docx import Document as DocxDocument
from PIL import Image
from PIL import ImageOps
import pytesseract
from pdf2image import convert_from_path
from openpyxl import load_workbook


ASK_CONVERSATION_SESSION_KEY = "ask_documents_conversation"


def get_session_user(request):
    return request.session.get("user", {})


def get_actor_name(request):
    user = get_session_user(request)
    return (
        user.get("name")
        or user.get("email")
        or user.get("preferred_username")
        or ""
    )


def get_actor_email(request):
    user = get_session_user(request)
    return user.get("email") or user.get("preferred_username") or ""


def get_document_audit_metadata(document):
    return {
        "document_type": document.document_type,
        "document_subtype": document.document_subtype,
        "department": document.department,
        "author": document.author,
        "tags": document.tags,
        "ocr_language": document.ocr_language,
        "file": document.file.name if document.file else "",
    }


def get_accessible_documents_queryset(request):
    if is_viewer(request):
        return Document.objects.all()

    return Document.objects.none()


def serialize_ask_turn(question, rag_result):
    citations = []

    for citation in rag_result.get("citations") or []:
        document = citation["document"]
        citations.append({
            "document_id": document.id,
            "chunk_text": (citation.get("chunk_text") or "")[:500],
            "score": citation.get("score") or 0,
        })

    return {
        "question": question,
        "answer": rag_result.get("answer") or "",
        "citations": citations,
    }


def hydrate_ask_conversation(stored_turns, accessible_documents):
    document_ids = {
        citation.get("document_id")
        for turn in stored_turns
        for citation in (turn.get("citations") or [])
        if citation.get("document_id") is not None
    }
    documents_by_id = accessible_documents.in_bulk(document_ids)
    conversation = []

    for turn in stored_turns:
        citations = []
        stored_citations = turn.get("citations") or []
        for stored_citation in stored_citations:
            document = documents_by_id.get(stored_citation.get("document_id"))
            if document is None:
                continue
            citations.append({
                "citation_id": len(citations) + 1,
                "document": document,
                "chunk_text": stored_citation.get("chunk_text") or "",
                "score": stored_citation.get("score") or 0,
            })

        if len(citations) != len(stored_citations):
            continue

        conversation.append({
            "question": (turn.get("question") or "").strip(),
            "answer": (turn.get("answer") or "").strip(),
            "citations": citations,
        })

    return conversation


def record_audit_event(request, document, action, metadata=None):
    AuditEvent.objects.create(
        document=document if document.pk else None,
        document_name=document.file.name if document.file else "",
        action=action,
        actor_name=get_actor_name(request),
        actor_email=get_actor_email(request),
        metadata=metadata or get_document_audit_metadata(document),
    )


def invalidate_metadata_quality_review(document):
    MetadataQualityReview.objects.filter(document_id=document.id).delete()


def get_ai_metadata_source_text(document):
    text = (document.extracted_text or "").strip()

    if text.startswith("TEXT_EXTRACTION_FAILED:"):
        return ""

    return text


def store_ai_metadata_suggestions(document):
    provider = settings.AI_METADATA_PROVIDER
    document.ai_suggestion_status = Document.AI_STATUS_PENDING
    document.ai_metadata_provider = provider
    document.ai_error = ""
    document.ai_explanation = {}
    document.save(
        update_fields=[
            "ai_suggestion_status",
            "ai_metadata_provider",
            "ai_error",
            "ai_explanation",
        ]
    )

    suggestions = suggest_metadata(
        get_ai_metadata_source_text(document)
    )

    document.ai_document_type = suggestions["document_type"]
    document.ai_department = suggestions["department"]
    document.ai_tags = suggestions["tags"]
    document.ai_summary = suggestions["summary"]
    document.ai_explanation = suggestions.get("explanation") or {}
    document.ai_suggestion_status = Document.AI_STATUS_SUGGESTED
    document.ai_suggested_at = timezone.now()
    document.ai_error = ""
    document.save(
        update_fields=[
            "ai_document_type",
            "ai_department",
            "ai_tags",
            "ai_summary",
            "ai_explanation",
            "ai_metadata_provider",
            "ai_suggestion_status",
            "ai_suggested_at",
            "ai_error",
        ]
    )


def try_store_ai_metadata_suggestions(document):
    if not settings.AUTO_AI_METADATA_ON_UPLOAD:
        return False

    try:
        store_ai_metadata_suggestions(document)
    except MetadataSuggestionError as error:
        document.ai_suggestion_status = Document.AI_STATUS_FAILED
        document.ai_error = str(error)
        document.save(
            update_fields=[
                "ai_suggestion_status",
                "ai_error",
            ]
        )
        return False

    return True


def try_rebuild_document_embeddings(request, document):
    if not get_ai_metadata_source_text(document):
        return False

    try:
        rebuild_document_embeddings(document)
    except EmbeddingError:
        messages.warning(
            request,
            "Document uploaded, but search embeddings could not be generated."
        )
        return False

    return True


def build_mcp_indexing_payload(document, replace_existing_chunks=True):
    return {
        "document": {
            "document_id": document.id,
            "file_name": document.file.name if document.file else "",
        },
        "options": {
            "replace_existing_chunks": replace_existing_chunks,
            "index_document_metadata": settings.OPENSEARCH_INDEX_ON_SAVE,
            "index_chunks": settings.OPENSEARCH_INDEX_ON_SAVE,
        },
        "trace": {
            "request_id": str(uuid4()),
            "source": "django-view",
            "actor": "django-backend",
        },
    }


def try_mcp_index_document(request, document, replace_existing_chunks=True):
    try:
        response = mcp_index_document(
            build_mcp_indexing_payload(
                document,
                replace_existing_chunks=replace_existing_chunks,
            )
        )
    except Exception:
        messages.warning(
            request,
            "Document saved, but MCP indexing could not be completed."
        )
        return False

    if response.get("status") != "ok":
        error = response.get("error") or {}
        code = error.get("code") or "internal_error"
        messages.warning(
            request,
            f"Document saved, but MCP indexing failed: {code}."
        )
        return False

    return True


def try_index_document_for_search(request, document):
    if settings.MCP_INDEXING_ENABLED:
        return try_mcp_index_document(
            request,
            document,
            replace_existing_chunks=True,
        )

    rebuilt = try_rebuild_document_embeddings(request, document)
    indexed = try_reindex_document(request, document)
    return rebuilt or indexed


def try_reindex_document(request, document):
    if not settings.OPENSEARCH_INDEX_ON_SAVE:
        return False

    if settings.MCP_INDEXING_ENABLED:
        return try_mcp_index_document(
            request,
            document,
            replace_existing_chunks=False,
        )

    try:
        reindex_document(document, create_indexes=True)
    except OpenSearchIndexingError:
        messages.warning(
            request,
            "Document saved, but OpenSearch indexing could not be completed."
        )
        return False

    return True


def try_delete_indexed_document(request, document_id):
    if not settings.OPENSEARCH_INDEX_ON_SAVE:
        return False

    try:
        delete_indexed_document(document_id)
    except OpenSearchIndexingError:
        messages.warning(
            request,
            "Document deleted, but OpenSearch cleanup could not be completed."
        )
        return False

    return True


def healthz(request):
    return HttpResponse("ok", content_type="text/plain")


def get_temp_upload_path(filename):
    temp_dir = os.path.abspath(os.path.join(settings.MEDIA_ROOT, "temp"))
    file_path = os.path.abspath(os.path.join(temp_dir, os.path.basename(filename)))

    if not file_path.startswith(temp_dir + os.sep):
        raise Http404("Temporary scanned file not found")

    return file_path


def get_ocr_language_choices():
    return Document.OCR_LANGUAGE_CHOICES


def get_ocr_language_label(ocr_language):
    return dict(Document.OCR_LANGUAGE_CHOICES).get(ocr_language, "English")


def normalize_ocr_language(ocr_language):
    allowed_languages = {
        value for value, label in Document.OCR_LANGUAGE_CHOICES
    }

    if ocr_language in allowed_languages:
        return ocr_language

    return Document.OCR_LANGUAGE_ENGLISH


def get_tesseract_language(ocr_language):
    language_map = {
        Document.OCR_LANGUAGE_ENGLISH: "eng",
        Document.OCR_LANGUAGE_HINDI: "hin",
        Document.OCR_LANGUAGE_URDU: "urd",
    }

    return language_map.get(
        normalize_ocr_language(ocr_language),
        "eng"
    )


def get_tesseract_config(ocr_language):
    if normalize_ocr_language(ocr_language) == Document.OCR_LANGUAGE_URDU:
        return "--psm 6 -c preserve_interword_spaces=1"

    return "--psm 6"


def preprocess_image_for_ocr(image):
    image = ImageOps.exif_transpose(image)
    image = image.convert("L")

    width, height = image.size

    if max(width, height) < 1800:
        image = image.resize((width * 2, height * 2))

    image = ImageOps.autocontrast(image)
    return image.point(lambda pixel: 255 if pixel > 175 else 0)


def login(request):
    return oauth.okta.authorize_redirect(
        request,
        settings.OKTA_CALLBACK_URL
    )


def normalize_okta_groups(raw_groups):
    group_names = []

    for group in raw_groups or []:
        if isinstance(group, str):
            group_names.append(group)

        elif isinstance(group, dict):
            name = group.get("profile", {}).get("name") or group.get("name")

            if name:
                group_names.append(name)

    return group_names


def oidc_callback(request):
    token = oauth.okta.authorize_access_token(request)

    userinfo = token.get("userinfo")

    if not userinfo:
        userinfo = oauth.okta.userinfo(token=token)

    decoded_token = jwt.decode(
        token["id_token"],
        options={"verify_signature": False}
    )

    group_names = []

    for claim_name in ("django_groups", "groups"):
        group_names.extend(normalize_okta_groups(decoded_token.get(claim_name)))

    group_names.extend(normalize_okta_groups(userinfo.get("groups")))
    group_names = sorted(set(group_names))

    request.session["user"] = {
        "sub": userinfo.get("sub"),
        "name": userinfo.get("name"),
        "email": userinfo.get("email"),
        "preferred_username": userinfo.get("preferred_username"),
        "groups": group_names,
    }

    return redirect("/")


def logout(request):
    django_logout(request)
    request.session.flush()

    return redirect(settings.OKTA_LOGOUT_REDIRECT_URL)


@okta_role_required(is_viewer)
def profile(request):
    return render(request, "profile.html", {
        "user": request.session["user"]
    })


@okta_role_required(is_loader)
def upload_document(request):
    if request.method == "POST":
        form = DocumentForm(request.POST, request.FILES)

        if form.is_valid():
            document = form.save(commit=False)
            document.ocr_language = normalize_ocr_language(
                document.ocr_language
            )

            if not document.author:
                document.author = (
                    request.session["user"].get("name")
                    or request.session["user"].get("email")
                )

            document.save()

            try:
                file_path = document.file.path
                document.extracted_text = extract_text_from_file(
                    file_path,
                    document.ocr_language
                )
                document.save()

            except Exception as e:
                document.extracted_text = f"TEXT_EXTRACTION_FAILED: {str(e)}"
                document.save()

            try_index_document_for_search(request, document)

            record_audit_event(request, document, AuditEvent.ACTION_UPLOAD)

            if try_store_ai_metadata_suggestions(document):
                document.refresh_from_db()
                try_reindex_document(request, document)
                messages.success(
                    request,
                    "AI metadata suggestions are ready for review."
                )
            elif document.ai_suggestion_status == Document.AI_STATUS_FAILED:
                document.refresh_from_db()
                try_reindex_document(request, document)
                messages.warning(
                    request,
                    "Document uploaded, but AI metadata suggestions failed."
                )

            return redirect("edit_document_metadata", document_id=document.id)

    else:
        form = DocumentForm()

    return render(request, "upload.html", {"form": form})

@okta_role_required(is_loader)
def upload_scanned_image(request):

    if "user" not in request.session:
        return redirect("/login/")

    if request.method == "POST":

        uploaded_file = request.FILES.get("file")
        ocr_language = normalize_ocr_language(
            request.POST.get("ocr_language")
        )

        if not uploaded_file:
            return render_upload_scanned(
                request,
                "Please select a file.",
                ocr_language
            )

        try:
            validate_uploaded_file(
                uploaded_file,
                settings.ALLOWED_SCANNED_IMAGE_EXTENSIONS
            )
        except ValidationError as error:
            return render_upload_scanned(
                request,
                error.messages[0],
                ocr_language
            )

        temp_dir = os.path.join(settings.MEDIA_ROOT, "temp")
        os.makedirs(temp_dir, exist_ok=True)

        fs = FileSystemStorage(
            location=temp_dir
        )

        temp_filename = fs.save(uploaded_file.name, uploaded_file)
        temp_file_path = os.path.join(temp_dir, temp_filename)
        temp_file_url = reverse(
            "temp_scanned_preview",
            kwargs={"filename": temp_filename}
        )

        ocr_text = extract_text_from_file(temp_file_path, ocr_language)
        metadata = extract_metadata_from_ocr(ocr_text)

        return render(request, "ocr_review.html", {
            "temp_filename": temp_filename,
            "temp_file_url": temp_file_url,
            "original_filename": uploaded_file.name,
            "ocr_language": ocr_language,
            "ocr_language_label": get_ocr_language_label(ocr_language),
            "ocr_text": ocr_text,
            "document_type": metadata.get("document_type", ""),
            "document_subtype": metadata.get("document_subtype", ""),
            
        })

    return render_upload_scanned(request)


def render_upload_scanned(
    request,
    error=None,
    selected_ocr_language=Document.OCR_LANGUAGE_ENGLISH
):
    return render(request, "upload_scanned.html", {
        "error": error,
        "ocr_language_choices": get_ocr_language_choices(),
        "selected_ocr_language": normalize_ocr_language(
            selected_ocr_language
        ),
    })


@okta_role_required(is_loader)
def temp_scanned_preview(request, filename):
    file_path = get_temp_upload_path(filename)

    if not os.path.exists(file_path):
        raise Http404("Temporary scanned file not found")

    return FileResponse(
        open(file_path, "rb"),
        as_attachment=False,
        filename=os.path.basename(file_path)
    )


@okta_role_required(is_loader)
def confirm_document(request):

    if "user" not in request.session:
        return redirect("/login/")

    if request.method == "POST":

        temp_filename = request.POST.get("temp_filename")
        document_type = request.POST.get("document_type", "")
        document_subtype = request.POST.get("document_subtype", "")
        ocr_language = normalize_ocr_language(
            request.POST.get("ocr_language")
        )
        ocr_text = request.POST.get("ocr_text", "")

        if not temp_filename:
            raise Http404("Temporary scanned file not found")

        temp_file_path = get_temp_upload_path(temp_filename)

        if not os.path.exists(temp_file_path):
            raise Http404("Temporary scanned file not found")

        document = Document()
        document.document_type = document_type
        document.document_subtype = document_subtype
        document.ocr_language = ocr_language
        document.extracted_text = ocr_text

        if not document.author:
            document.author = (
                request.session["user"].get("name")
                or request.session["user"].get("email")
            )

        with open(temp_file_path, "rb") as f:
            document.file.save(
                os.path.basename(temp_file_path),
                File(f),
                save=False
            )

        document.save()
        try_index_document_for_search(request, document)
        record_audit_event(request, document, AuditEvent.ACTION_UPLOAD)

        if try_store_ai_metadata_suggestions(document):
            document.refresh_from_db()
            try_reindex_document(request, document)
            messages.success(
                request,
                "AI metadata suggestions are ready for review."
            )
        elif document.ai_suggestion_status == Document.AI_STATUS_FAILED:
            document.refresh_from_db()
            try_reindex_document(request, document)
            messages.warning(
                request,
                "Document saved, but AI metadata suggestions failed."
            )

        try:
            os.remove(temp_file_path)
        except Exception:
            pass

        return redirect("edit_document_metadata", document_id=document.id)

    return redirect("upload_scanned_image")


@okta_role_required(is_viewer)
def search_documents(request):
    documents = Document.objects.all().order_by("-uploaded_at")

    document_type = request.GET.get("document_type")
    document_subtype = request.GET.get("document_subtype")
    department = request.GET.get("department")
    author = request.GET.get("author")
    tags = request.GET.get("tags")
    content = request.GET.get("content")

    if document_type:
        documents = documents.filter(document_type__icontains=document_type)
    if document_subtype:
        documents = documents.filter(document_subtype__icontains=document_subtype)    

    if department:
        documents = documents.filter(department__icontains=department)

    if author:
        documents = documents.filter(author__icontains=author)

    if tags:
        documents = documents.filter(tags__icontains=tags)

    if content:
        documents = documents.filter(
            Q(extracted_text__icontains=content)
        )

    paginator = Paginator(documents, settings.DOCUMENTS_PER_PAGE)
    page_obj = paginator.get_page(request.GET.get("page"))
    query_params = request.GET.copy()
    query_params.pop("page", None)

    return render(request, "search.html", {
        "documents": page_obj,
        "page_obj": page_obj,
        "query_string": query_params.urlencode(),
        "total_documents": paginator.count,
    })


@okta_role_required(is_viewer)
def ai_search(request):
    query = (request.GET.get("q") or "").strip()
    results = []
    has_embeddings = (
        DocumentChunk.objects
        .exclude(embedding=[])
        .filter(embedding_model=settings.BEDROCK_EMBED_MODEL_ID)
        .exists()
    )

    if query and has_embeddings:
        try:
            results = search_documents_by_meaning(query)
        except EmbeddingError as error:
            messages.warning(
                request,
                f"AI search could not be completed: {error}"
            )

    return render(request, "ai_search.html", {
        "query": query,
        "results": results,
        "has_embeddings": has_embeddings,
    })


@okta_role_required(is_viewer)
def ask_documents(request):
    if request.GET.get("new") == "1":
        request.session.pop(ASK_CONVERSATION_SESSION_KEY, None)
        return redirect("ask_documents")

    question = (request.GET.get("q") or "").strip()
    rag_result = None
    accessible_documents = get_accessible_documents_queryset(request)
    stored_conversation = request.session.get(
        ASK_CONVERSATION_SESSION_KEY,
        [],
    )
    conversation = hydrate_ask_conversation(
        stored_conversation,
        accessible_documents,
    )
    has_embeddings = (
        DocumentChunk.objects
        .exclude(embedding=[])
        .filter(embedding_model=settings.BEDROCK_EMBED_MODEL_ID)
        .filter(document__in=accessible_documents)
        .exists()
    )

    if question and has_embeddings:
        try:
            rag_result = answer_question(
                question,
                documents_queryset=accessible_documents,
                conversation_history=conversation,
            )
            stored_conversation.append(
                serialize_ask_turn(question, rag_result)
            )
            stored_conversation = stored_conversation[
                -settings.AI_RAG_CONVERSATION_MAX_TURNS:
            ]
            request.session[ASK_CONVERSATION_SESSION_KEY] = stored_conversation
            request.session.modified = True
        except (EmbeddingError, RAGError) as error:
            messages.warning(
                request,
                f"Document Q&A could not be completed: {error}"
            )

    conversation = hydrate_ask_conversation(
        stored_conversation,
        accessible_documents,
    )

    return render(request, "ask_documents.html", {
        "question": question,
        "rag_result": rag_result,
        "conversation": conversation,
        "has_embeddings": has_embeddings,
    })


@okta_role_required(is_admin)
def audit_events(request):
    events = AuditEvent.objects.select_related("document").all()
    paginator = Paginator(events, settings.DOCUMENTS_PER_PAGE)
    page_obj = paginator.get_page(request.GET.get("page"))

    return render(request, "audit_events.html", {
        "events": page_obj,
        "page_obj": page_obj,
        "total_events": paginator.count,
    })


@okta_role_required(is_loader)
def metadata_quality_dashboard(request):
    active_filter = (request.GET.get("filter") or "all").strip().lower()
    valid_filters = {
        "all",
        "critical",
        "needs_review",
        "good",
        "excellent",
        "unreviewed",
        "failed",
    }
    if active_filter not in valid_filters:
        active_filter = "all"

    documents = list(Document.objects.order_by("-uploaded_at"))
    reviews = {
        review.document_id: review
        for review in MetadataQualityReview.objects.all()
    }
    complete_reviews = [
        review
        for review in reviews.values()
        if review.status == MetadataQualityReview.STATUS_COMPLETE
    ]

    rows = []
    for document in documents:
        review = reviews.get(document.id)
        include = active_filter == "all"

        if active_filter == "unreviewed":
            include = review is None
        elif active_filter == "failed":
            include = (
                review is not None
                and review.status == MetadataQualityReview.STATUS_FAILED
            )
        elif active_filter in {
            "critical",
            "needs_review",
            "good",
            "excellent",
        }:
            include = (
                review is not None
                and review.status == MetadataQualityReview.STATUS_COMPLETE
                and review.quality_level == active_filter
            )

        if include:
            rows.append({
                "document": document,
                "review": review,
            })

    average_score = (
        round(
            sum(review.quality_score for review in complete_reviews)
            / len(complete_reviews)
        )
        if complete_reviews
        else None
    )
    level_counts = {
        level: sum(
            1
            for review in complete_reviews
            if review.quality_level == level
        )
        for level in (
            MetadataQualityReview.LEVEL_EXCELLENT,
            MetadataQualityReview.LEVEL_GOOD,
            MetadataQualityReview.LEVEL_NEEDS_REVIEW,
            MetadataQualityReview.LEVEL_CRITICAL,
        )
    }

    paginator = Paginator(rows, settings.DOCUMENTS_PER_PAGE)
    page_obj = paginator.get_page(request.GET.get("page"))

    return render(request, "metadata_quality_dashboard.html", {
        "rows": page_obj,
        "page_obj": page_obj,
        "active_filter": active_filter,
        "total_documents": len(documents),
        "reviewed_count": len(complete_reviews),
        "unreviewed_count": sum(
            1 for document in documents if document.id not in reviews
        ),
        "failed_count": sum(
            1
            for review in reviews.values()
            if review.status == MetadataQualityReview.STATUS_FAILED
        ),
        "average_score": average_score,
        "level_counts": level_counts,
        "batch_limit": settings.AI_METADATA_QUALITY_BATCH_LIMIT,
    })


@okta_role_required(is_loader)
def metadata_quality_detail(request, document_id):
    try:
        document = Document.objects.get(id=document_id)
    except Document.DoesNotExist:
        raise Http404("Document not found")

    try:
        review = document.metadata_quality_review
    except MetadataQualityReview.DoesNotExist:
        review = None

    return render(request, "metadata_quality_detail.html", {
        "document": document,
        "review": review,
    })


def perform_metadata_quality_review(request, document):
    try:
        result = review_document_metadata(document)
    except MetadataQualityError as error:
        MetadataQualityReview.objects.update_or_create(
            document=document,
            defaults={
                "status": MetadataQualityReview.STATUS_FAILED,
                "quality_score": 0,
                "quality_level": "",
                "review_summary": "",
                "issues": [],
                "model_id": settings.BEDROCK_NOVA_MODEL_ID,
                "error": str(error),
            },
        )
        return None, str(error)

    review, _ = MetadataQualityReview.objects.update_or_create(
        document=document,
        defaults={
            "status": MetadataQualityReview.STATUS_COMPLETE,
            "quality_score": result["quality_score"],
            "quality_level": result["quality_level"],
            "review_summary": result["summary"],
            "issues": result["issues"],
            "model_id": settings.BEDROCK_NOVA_MODEL_ID,
            "error": "",
        },
    )
    record_audit_event(
        request,
        document,
        AuditEvent.ACTION_EDIT,
        {
            "source": "ai_metadata_quality_review",
            "quality_score": review.quality_score,
            "quality_level": review.quality_level,
            "issue_count": review.issue_count,
            "model_id": review.model_id,
        },
    )
    return review, ""


@okta_role_required(is_loader)
def run_metadata_quality_review(request, document_id):
    if request.method != "POST":
        return redirect("metadata_quality_detail", document_id=document_id)

    try:
        document = Document.objects.get(id=document_id)
    except Document.DoesNotExist:
        raise Http404("Document not found")

    review, error = perform_metadata_quality_review(request, document)
    if error:
        messages.error(request, f"Metadata quality review failed: {error}")
        return redirect("metadata_quality_detail", document_id=document.id)

    messages.success(
        request,
        f"AI metadata quality review completed with score {review.quality_score}.",
    )
    return redirect("metadata_quality_detail", document_id=document.id)


@okta_role_required(is_loader)
def run_selected_metadata_quality_reviews(request):
    if request.method != "POST":
        return redirect("metadata_quality_dashboard")

    selected_ids = []
    for value in request.POST.getlist("document_ids"):
        try:
            selected_ids.append(int(value))
        except (TypeError, ValueError):
            continue

    selected_ids = list(dict.fromkeys(selected_ids))
    if not selected_ids:
        messages.warning(request, "Select at least one document to review.")
        return redirect("metadata_quality_dashboard")

    batch_limit = settings.AI_METADATA_QUALITY_BATCH_LIMIT
    if len(selected_ids) > batch_limit:
        messages.warning(
            request,
            f"Review up to {batch_limit} documents at a time.",
        )
        selected_ids = selected_ids[:batch_limit]

    documents_by_id = Document.objects.in_bulk(selected_ids)
    succeeded = 0
    failed = 0

    for document_id in selected_ids:
        document = documents_by_id.get(document_id)
        if document is None:
            failed += 1
            continue

        review, error = perform_metadata_quality_review(request, document)
        if review is not None and not error:
            succeeded += 1
        else:
            failed += 1

    if succeeded:
        messages.success(
            request,
            f"AI metadata quality review completed for {succeeded} "
            f"document{'s' if succeeded != 1 else ''}.",
        )
    if failed:
        messages.warning(
            request,
            f"{failed} document review{'s' if failed != 1 else ''} failed.",
        )

    return redirect("metadata_quality_dashboard")


@okta_role_required(is_viewer)
def secure_document_view(request, document_id):
    try:
        document = Document.objects.get(id=document_id)

    except Document.DoesNotExist:
        raise Http404("Document not found")

    file_path = document.file.path

    if not os.path.exists(file_path):
        raise Http404("File not found")

    return FileResponse(
        open(file_path, "rb"),
        as_attachment=False,
        filename=os.path.basename(file_path)
    )


@okta_role_required(is_loader)
def edit_document_metadata(request, document_id):
    try:
        document = Document.objects.get(id=document_id)

    except Document.DoesNotExist:
        raise Http404("Document not found")

    if request.method == "POST":
        form = DocumentMetadataForm(request.POST, instance=document)

        if form.is_valid():
            before_metadata = get_document_audit_metadata(document)
            form.save()
            document.refresh_from_db()
            invalidate_metadata_quality_review(document)
            after_metadata = get_document_audit_metadata(document)
            record_audit_event(
                request,
                document,
                AuditEvent.ACTION_EDIT,
                {
                    "before": before_metadata,
                    "after": after_metadata,
                },
            )
            try_reindex_document(request, document)
            return redirect("search")

    else:
        form = DocumentMetadataForm(instance=document)

    file_extension = os.path.splitext(document.file.name)[1].lower()
    active_provider = settings.AI_METADATA_PROVIDER
    active_provider_label = dict(
        Document.AI_METADATA_PROVIDER_CHOICES
    ).get(active_provider, active_provider)

    return render(request, "edit_metadata.html", {
        "document": document,
        "form": form,
        "active_ai_metadata_provider_label": active_provider_label,
        "can_preview_inline": file_extension in [
            ".png",
            ".jpg",
            ".jpeg",
            ".tiff",
            ".bmp",
        ],
    })


@okta_role_required(is_loader)
def generate_ai_metadata(request, document_id):
    if request.method != "POST":
        return redirect("edit_document_metadata", document_id=document_id)

    try:
        document = Document.objects.get(id=document_id)
    except Document.DoesNotExist:
        raise Http404("Document not found")

    try:
        store_ai_metadata_suggestions(document)
    except MetadataSuggestionError as error:
        document.ai_suggestion_status = Document.AI_STATUS_FAILED
        document.ai_error = str(error)
        document.save(
            update_fields=[
                "ai_suggestion_status",
                "ai_error",
            ]
        )
        document.refresh_from_db()
        try_reindex_document(request, document)
        messages.error(request, str(error))
        return redirect("edit_document_metadata", document_id=document.id)

    document.refresh_from_db()
    try_reindex_document(request, document)
    messages.success(request, "AI metadata suggestions generated.")

    return redirect("edit_document_metadata", document_id=document.id)


@okta_role_required(is_loader)
def accept_ai_metadata(request, document_id):
    if request.method != "POST":
        return redirect("edit_document_metadata", document_id=document_id)

    try:
        document = Document.objects.get(id=document_id)
    except Document.DoesNotExist:
        raise Http404("Document not found")

    before_metadata = get_document_audit_metadata(document)

    if document.ai_document_type:
        document.document_type = document.ai_document_type
    if document.ai_department:
        document.department = document.ai_department
    if document.ai_tags:
        document.tags = document.ai_tags
    if document.ai_summary:
        document.description = document.ai_summary

    document.ai_suggestion_status = Document.AI_STATUS_ACCEPTED
    document.save()
    document.refresh_from_db()
    invalidate_metadata_quality_review(document)

    record_audit_event(
        request,
        document,
        AuditEvent.ACTION_EDIT,
        {
            "source": "ai_metadata_accept",
            "ai_suggestion": {
                "document_type": document.ai_document_type,
                "department": document.ai_department,
                "tags": document.ai_tags,
                "summary": document.ai_summary,
                "explanation": document.ai_explanation,
            },
            "before": before_metadata,
            "after": get_document_audit_metadata(document),
        },
    )
    messages.success(request, "AI suggestions accepted into metadata.")
    try_reindex_document(request, document)

    return redirect("edit_document_metadata", document_id=document.id)


@okta_role_required(is_loader)
def reject_ai_metadata(request, document_id):
    if request.method != "POST":
        return redirect("edit_document_metadata", document_id=document_id)

    try:
        document = Document.objects.get(id=document_id)
    except Document.DoesNotExist:
        raise Http404("Document not found")

    rejected_suggestion = {
        "document_type": document.ai_document_type,
        "department": document.ai_department,
        "tags": document.ai_tags,
        "summary": document.ai_summary,
        "explanation": document.ai_explanation,
    }
    document.ai_suggestion_status = Document.AI_STATUS_REJECTED
    document.save(update_fields=["ai_suggestion_status"])
    document.refresh_from_db()
    record_audit_event(
        request,
        document,
        AuditEvent.ACTION_EDIT,
        {
            "source": "ai_metadata_reject",
            "ai_suggestion": rejected_suggestion,
        },
    )
    try_reindex_document(request, document)
    messages.info(request, "AI suggestions rejected.")

    return redirect("edit_document_metadata", document_id=document.id)


@okta_role_required(is_admin)
def delete_document(request, document_id):
    try:
        document = Document.objects.get(id=document_id)

    except Document.DoesNotExist:
        raise Http404("Document not found")

    if request.method == "POST":
        file_path = document.file.path
        audit_metadata = get_document_audit_metadata(document)

        if file_path and os.path.exists(file_path):
            os.remove(file_path)

        record_audit_event(
            request,
            document,
            AuditEvent.ACTION_DELETE,
            audit_metadata,
        )
        document_id = document.id
        document.delete()
        try_delete_indexed_document(request, document_id)

        return redirect("/")

    return render(request, "delete_document.html", {
        "document": document,
    })

def extract_metadata_from_ocr(ocr_text):
    metadata = {
        "document_type": "",
        "document_subtype": "",
    }

    for line in ocr_text.splitlines():
        line_lower = line.lower()

        if "document type" in line_lower:
            parts = line.split(":", 1)
            if len(parts) == 2:
                metadata["document_type"] = parts[1].strip()
            else:
                metadata["document_type"] = line.strip()
        elif "document subtype" in line_lower:
            parts = line.split(":", 1)
            if len(parts) == 2:
                metadata["document_subtype"] = parts[1].strip()        

    return metadata


def extract_text_from_file(file_path, ocr_language=Document.OCR_LANGUAGE_ENGLISH):
    text = ""
    ext = os.path.splitext(file_path)[1].lower()
    tesseract_language = get_tesseract_language(ocr_language)
    tesseract_config = get_tesseract_config(ocr_language)

    if ext == ".pdf":
        reader = PdfReader(file_path)

        for page in reader.pages:
            text += page.extract_text() or ""
            text += "\n"

        if not text.strip():
            images = convert_from_path(file_path)

            for image in images:
                image = preprocess_image_for_ocr(image)
                text += pytesseract.image_to_string(
                    image,
                    lang=tesseract_language,
                    config=tesseract_config
                )
                text += "\n"

    elif ext == ".docx":
        doc = DocxDocument(file_path)

        for para in doc.paragraphs:
            text += para.text + "\n"

    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()

    elif ext == ".xlsx":
        workbook = load_workbook(file_path, data_only=True)

        for sheet in workbook.worksheets:
            text += f"\n--- Sheet: {sheet.title} ---\n"

            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join(
                    str(cell) for cell in row if cell is not None
                )
                text += row_text + "\n"

    elif ext in [".png", ".jpg", ".jpeg", ".tiff", ".bmp"]:
        image = Image.open(file_path)

        image = preprocess_image_for_ocr(image)

        text = pytesseract.image_to_string(
            image,
            lang=tesseract_language,
            config=tesseract_config
        )

    return text
